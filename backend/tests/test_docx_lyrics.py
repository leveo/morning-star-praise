# SPDX-License-Identifier: GPL-3.0-or-later
import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.main import app


client = TestClient(app)


def _make_source(*, include_chorus: bool = True, omit_english_verse_2: bool = False) -> bytes:
    doc = Document()
    paragraphs = [
        "450 主啊！我今來",
        "1",
        "我心何等歡喜，",
        "因聽救主說道;",
        "我為你罪流出寶血, ",
        "使你同得榮耀。",
    ]
    if include_chorus:
        paragraphs.extend([
            "副歌：",
            "主啊！我今來，我今來就你,",
            "求主洗淨我罪愆, ",
            "洗淨在寶血裏。",
        ])
    paragraphs.extend([
        "2",
        "我本軟弱可憐, ",
        "行善毫無能力;",
        "惟主能顯救恩奇功,",
        "救我脫離罪權。",
        "3",
        "求主賜我聖靈, ",
        "充滿在我的心;",
        "更求我主恩上加恩, ",
        "變成救主容形。",
        "450 I am Coming, Lord",
        "I hear Thy welcome voice, ",
        "That calls me, Lord, to Thee,",
        "For cleansing in Thy precious blood",
        "That flowed on Calvary.",
    ])
    if include_chorus:
        paragraphs.extend([
            "Refrain:",
            "I am coming, Lord!",
            "Coming now to Thee! ",
            "Wash me, cleanse me in the blood",
            "That flowed on Calvary.",
        ])
    if not omit_english_verse_2:
        paragraphs.extend([
            "2",
            "Though coming weak and vile, ",
            "Thou do my strength assure; ",
            "Thou do my vileness fully cleanse, ",
            "Till spotless all and pure.",
        ])
    paragraphs.extend([
        "3",
        "Tis Jesus calls me on",
        "To perfect faith and love,",
        "To perfect hope, and peace, and trust,",
        "For earth and heaven above.",
    ])
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _import_source(content: bytes | None = None, filename: str = "450 主啊！我今來 I AM COMING, LORD.docx"):
    return client.post(
        "/api/lyrics/import-docx",
        files={"file": (filename, content or _make_source(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )


def test_import_docx_expands_chorus_after_every_verse():
    response = _import_source()
    assert response.status_code == 200
    data = response.json()
    assert data["song_number"] == "450"
    assert data["title_zh"] == "主啊！我今來"
    assert data["title_en"] == "I am Coming, Lord"
    assert data["sequence"] == [
        "verse-1", "chorus", "verse-2", "chorus", "verse-3", "chorus"
    ]
    assert data["has_blocking_errors"] is False
    assert data["combined_lyrics"].count("I am coming, Lord!") == 3
    assert "副歌" not in data["combined_lyrics"]
    assert "Refrain" not in data["combined_lyrics"]


def test_import_docx_without_chorus_keeps_verses_only():
    response = _import_source(_make_source(include_chorus=False))
    assert response.status_code == 200
    data = response.json()
    assert data["sequence"] == ["verse-1", "verse-2", "verse-3"]
    assert [section["kind"] for section in data["sections"]] == ["verse"] * 3


def test_import_docx_reports_missing_language_section():
    response = _import_source(_make_source(omit_english_verse_2=True))
    assert response.status_code == 200
    data = response.json()
    assert data["has_blocking_errors"] is True
    assert any(
        warning["code"] == "missing_section_en" and warning["section_id"] == "verse-2"
        for warning in data["warnings"]
    )


def test_import_rejects_generated_v2_and_non_docx():
    generated = _import_source(filename="450 song_v2.docx")
    assert generated.status_code == 400
    assert "_v2" in generated.json()["detail"]

    wrong_type = _import_source(filename="450 song.doc")
    assert wrong_type.status_code == 400
    assert ".docx" in wrong_type.json()["detail"]


def test_export_docx_has_required_order_fonts_and_filename():
    imported = _import_source().json()
    payload = {
        key: imported[key]
        for key in (
            "source_filename",
            "song_number",
            "title_zh",
            "title_en",
            "collection_zh",
            "collection_en",
            "sections",
        )
    }
    response = client.post("/api/lyrics/export-docx", json=payload)
    assert response.status_code == 200
    assert "filename*=UTF-8''450%20" in response.headers["content-disposition"]

    exported = Document(io.BytesIO(response.content))
    texts = [paragraph.text for paragraph in exported.paragraphs]
    assert texts.count("主啊！我今來，我今來就你,") == 3
    assert texts.count("I am coming, Lord!") == 3

    for paragraph in exported.paragraphs:
        if not paragraph.text:
            continue
        for run in paragraph.runs:
            assert run.font.size.pt == 12
            fonts = run._element.rPr.rFonts
            if any("\u4e00" <= char <= "\u9fff" for char in paragraph.text):
                assert fonts.get(qn("w:eastAsia")) == "宋体"
            else:
                assert fonts.get(qn("w:ascii")) == "Times New Roman"
                assert fonts.get(qn("w:hAnsi")) == "Times New Roman"

    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        font_table = archive.read("word/fontTable.xml").decode("utf-8")
    assert 'w:name="宋体"' in font_table
    assert 'w:altName w:val="Songti TC"' in font_table


def test_export_refuses_incomplete_sections():
    imported = _import_source(_make_source(omit_english_verse_2=True)).json()
    payload = {
        key: imported[key]
        for key in (
            "source_filename",
            "song_number",
            "title_zh",
            "title_en",
            "collection_zh",
            "collection_en",
            "sections",
        )
    }
    response = client.post("/api/lyrics/export-docx", json=payload)
    assert response.status_code == 422
    assert "第 2 节缺少英文歌词" in response.json()["detail"]
