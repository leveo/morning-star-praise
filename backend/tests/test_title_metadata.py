# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Leo Song
from pptx import Presentation

from app.services import ppt_service


def test_ppt_title_metadata_uses_separate_editable_text_boxes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    ppt_service._add_title_slide(
        slide,
        "主啊！我今来",
        "Author",
        "zh-hant",
        prs.slide_width,
        prs.slide_height,
        title_en="I Am Coming, Lord",
        collection_zh="教會聖詩 #450",
        collection_en="Hymn's for God's People",
    )

    text_boxes = [shape.text for shape in slide.shapes if shape.has_text_frame]
    assert "主啊！我今来" in text_boxes
    assert "I Am Coming, Lord" in text_boxes
    assert "教會聖詩 #450" in text_boxes
    assert "Hymn's for God's People" in text_boxes
    assert "Author" not in text_boxes


def test_ppt_title_falls_back_to_composer_without_collection_metadata():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    ppt_service._add_title_slide(
        slide,
        "Old Song",
        "Traditional",
        "en",
        prs.slide_width,
        prs.slide_height,
    )

    text_boxes = [shape.text for shape in slide.shapes if shape.has_text_frame]
    assert "Traditional" in text_boxes
