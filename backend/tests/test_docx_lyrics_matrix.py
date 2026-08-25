# SPDX-License-Identifier: GPL-3.0-or-later
"""Offline layout matrix for deterministic bilingual hymn parsing.

The text below is synthetic. The adjacent reference manifest records only
public URLs and structural fingerprints, so regression tests never redistribute
third-party hymn text.
"""

from __future__ import annotations

from dataclasses import dataclass
import io
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import pytest

from app.services.docx_lyrics_service import parse_docx
from app.services.docx_lyrics_parser import extract_fragments


@dataclass(frozen=True)
class Truth:
    number: int
    slug: str
    zh_title: str
    en_title: str
    verses: int
    chorus: bool
    amen: bool = False

    def zh(self, verse: int) -> list[str]:
        return [f"{self.zh_title}第{verse}節甲句", f"{self.zh_title}第{verse}節乙句"]

    def en(self, verse: int) -> list[str]:
        return [f"{self.en_title} verse {verse} line A", f"{self.en_title} verse {verse} line B"]

    @property
    def chorus_zh(self) -> list[str]:
        return [f"{self.zh_title}副歌甲句", f"{self.zh_title}副歌乙句"]

    @property
    def chorus_en(self) -> list[str]:
        return [f"{self.en_title} refrain line A", f"{self.en_title} refrain line B"]


TRUTHS = [
    Truth(450, "welcome", "主啊我今來", "Welcome Voice", 3, True, True),
    Truth(465, "story", "我愛傳講", "Tell the Story", 4, True),
    Truth(468, "zion", "信徒奮興", "O Zion Haste", 4, True),
    Truth(101, "grace", "奇異恩典", "Amazing Grace", 5, False, True),
    Truth(102, "holy", "聖哉聖哉", "Holy Holy Holy", 4, False),
    Truth(103, "assurance", "有福確據", "Blessed Assurance", 3, True),
    Truth(104, "glory", "榮耀歸神", "To God Be Glory", 3, True),
    Truth(105, "obey", "信靠順服", "Trust and Obey", 5, True),
    Truth(106, "loves", "耶穌愛我", "Jesus Loves Me", 3, True),
    Truth(107, "hail", "擁戴我主", "All Hail the Power", 4, False),
    Truth(108, "rock", "萬古磐石", "Rock of Ages", 4, False),
    Truth(109, "friend", "耶穌恩友", "What a Friend", 3, False),
]

FORMATS = [
    "blocks_plain",
    "blocks_inline_labels",
    "blocks_soft_breaks",
    "blocks_page_breaks",
    "blocks_hyperlinks",
    "blocks_tracked_changes",
    "stanza_interleaved",
    "line_interleaved",
    "table_columns",
    "table_stanzas",
    "automatic_numbering",
    "textbox",
    "refrain_references",
    "repeated_refrain",
]


def _save(document: Document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _add_hyperlink_text(paragraph, text: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "local")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_tracked_text(paragraph, text: str) -> None:
    deleted = OxmlElement("w:del")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "DELETED TEXT"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    inserted = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    inserted.append(run)
    paragraph._p.extend([deleted, inserted])


def _set_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend([level, number])
    properties.append(num_pr)


def _add_textbox(document: Document, lines: list[str]) -> None:
    paragraphs = "".join(
        f'<w:p><w:r><w:t xml:space="preserve">{line}</w:t></w:r></w:p>'
        for line in lines
    )
    xml = (
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        '<w:pict><v:shape><v:textbox>'
        f'<w:txbxContent>{paragraphs}</w:txbxContent>'
        '</v:textbox></v:shape></w:pict></w:r>'
    )
    document.add_paragraph()._p.append(parse_xml(xml))


def _block_lines(truth: Truth, language: str, *, references: bool = False) -> list[str]:
    result: list[str] = []
    for verse in range(1, truth.verses + 1):
        if not (language == "en" and verse == 1):
            result.append(f"Verse {verse}" if language == "en" else f"第{verse}節")
        result.extend(truth.en(verse) if language == "en" else truth.zh(verse))
        if truth.chorus and verse == 1:
            result.append("Refrain:" if language == "en" else "副歌：")
            result.extend(truth.chorus_en if language == "en" else truth.chorus_zh)
        elif truth.chorus and references:
            result.append("[Refrain]")
    if truth.amen:
        result.append("Amen" if language == "en" else "阿門。")
    return result


def _add_lines(document: Document, lines: list[str], mode: str) -> None:
    if mode == "soft":
        paragraph = document.add_paragraph()
        for index, line in enumerate(lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line)
        return
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        if mode == "hyperlink":
            _add_hyperlink_text(paragraph, line)
        elif mode == "tracked":
            _add_tracked_text(paragraph, line)
        else:
            paragraph.add_run(line)
            if mode == "page" and index and index % 4 == 0:
                paragraph.add_run().add_break(WD_BREAK.PAGE)


def _make_blocks(truth: Truth, variant: str) -> Document:
    document = Document()
    document.add_paragraph(f"{truth.number} {truth.zh_title}")
    zh = _block_lines(truth, "zh", references=variant == "refrain_references")
    en = _block_lines(truth, "en", references=variant == "refrain_references")
    if variant == "blocks_inline_labels":
        zh = [line.replace("第", "第", 1).replace("節", "節：", 1) if line.startswith("第") else line for line in zh]
        en = [f"{line}:" if line.startswith("Verse ") else line for line in en]
    mode = {
        "blocks_soft_breaks": "soft",
        "blocks_page_breaks": "page",
        "blocks_hyperlinks": "hyperlink",
        "blocks_tracked_changes": "tracked",
    }.get(variant, "plain")
    _add_lines(document, zh, mode)
    document.add_paragraph(f"{truth.number} {truth.en_title}")
    _add_lines(document, en, mode)
    document.add_paragraph("Copyright: synthetic offline parser fixture")
    return document


def _make_mixed(truth: Truth, line_interleaved: bool) -> Document:
    document = Document()
    document.add_paragraph(f"{truth.number} {truth.zh_title} {truth.en_title}")
    for verse in range(1, truth.verses + 1):
        document.add_paragraph(f"Verse {verse}:")
        zh, en = truth.zh(verse), truth.en(verse)
        lines = [item for pair in zip(zh, en) for item in pair] if line_interleaved else [*zh, *en]
        for line in lines:
            document.add_paragraph(line)
        if truth.chorus and verse == 1:
            document.add_paragraph("Refrain:")
            chorus = [item for pair in zip(truth.chorus_zh, truth.chorus_en) for item in pair]
            for line in chorus:
                document.add_paragraph(line)
    if truth.amen:
        document.add_paragraph("阿門。")
        document.add_paragraph("Amen")
    return document


def _make_table(truth: Truth, columns: bool) -> Document:
    document = Document()
    document.add_paragraph(f"{truth.number} {truth.zh_title} {truth.en_title}")
    if columns:
        table = document.add_table(rows=0, cols=2)
        for verse in range(1, truth.verses + 1):
            cells = table.add_row().cells
            cells[0].text = "\n".join([f"第{verse}節", *truth.zh(verse)])
            en_label = [] if verse == 1 else [f"Verse {verse}"]
            cells[1].text = "\n".join([*en_label, *truth.en(verse)])
        if truth.chorus:
            cells = table.add_row().cells
            cells[0].text = "\n".join(["副歌：", *truth.chorus_zh])
            cells[1].text = "\n".join(["Refrain:", *truth.chorus_en])
    else:
        table = document.add_table(rows=0, cols=1)
        for verse in range(1, truth.verses + 1):
            table.add_row().cells[0].text = "\n".join([
                f"Verse {verse}:", *truth.zh(verse), *truth.en(verse), "",
            ])
        if truth.chorus:
            table.add_row().cells[0].text = "\n".join([
                "Refrain:", *truth.chorus_zh, *truth.chorus_en,
            ])
    if truth.amen:
        document.add_paragraph("阿門。")
        document.add_paragraph("Amen")
    return document


def _make_automatic_numbering(truth: Truth) -> Document:
    document = Document()
    document.add_paragraph(f"{truth.number} {truth.zh_title}")
    for verse in range(1, truth.verses + 1):
        marker = document.add_paragraph(truth.zh(verse)[0])
        _set_numbering(marker, 91)
        document.add_paragraph(truth.zh(verse)[1])
        if truth.chorus and verse == 1:
            for line in ["副歌：", *truth.chorus_zh]:
                document.add_paragraph(line)
    document.add_paragraph(f"{truth.number} {truth.en_title}")
    for verse in range(1, truth.verses + 1):
        marker = document.add_paragraph(truth.en(verse)[0])
        _set_numbering(marker, 92)
        document.add_paragraph(truth.en(verse)[1])
        if truth.chorus and verse == 1:
            for line in ["Refrain:", *truth.chorus_en]:
                document.add_paragraph(line)
    if truth.amen:
        document.add_paragraph("Amen")
    return document


def _make_repeated(truth: Truth) -> Document:
    if not truth.chorus:
        return _make_mixed(truth, False)
    document = Document()
    document.add_paragraph(f"{truth.number} {truth.zh_title} {truth.en_title}")
    document.add_paragraph("")
    for verse in range(1, truth.verses + 1):
        for line in [*truth.zh(verse), *truth.en(verse), *truth.chorus_zh, *truth.chorus_en]:
            document.add_paragraph(line)
        if truth.amen and verse == truth.verses:
            document.add_paragraph("阿門。")
            document.add_paragraph("Amen")
        document.add_paragraph("")
    return document


def make_document(truth: Truth, variant: str) -> Document:
    if variant in {
        "blocks_plain", "blocks_inline_labels", "blocks_soft_breaks",
        "blocks_page_breaks", "blocks_hyperlinks", "blocks_tracked_changes",
        "refrain_references",
    }:
        return _make_blocks(truth, variant)
    if variant == "stanza_interleaved":
        return _make_mixed(truth, False)
    if variant == "line_interleaved":
        return _make_mixed(truth, True)
    if variant == "table_columns":
        return _make_table(truth, True)
    if variant == "table_stanzas":
        return _make_table(truth, False)
    if variant == "automatic_numbering":
        return _make_automatic_numbering(truth)
    if variant == "textbox":
        document = Document()
        document.add_paragraph(f"{truth.number} {truth.zh_title}")
        _add_textbox(document, _block_lines(truth, "zh"))
        document.add_paragraph(f"{truth.number} {truth.en_title}")
        _add_textbox(document, _block_lines(truth, "en"))
        return document
    if variant == "repeated_refrain":
        return _make_repeated(truth)
    raise AssertionError(variant)


def expected_sections(truth: Truth) -> list[dict]:
    result = [
        {"id": f"verse-{verse}", "zh": truth.zh(verse), "en": truth.en(verse)}
        for verse in range(1, truth.verses + 1)
    ]
    if truth.chorus:
        result.append({"id": "chorus", "zh": truth.chorus_zh, "en": truth.chorus_en})
    return result


@pytest.mark.parametrize("truth", TRUTHS, ids=lambda value: value.slug)
@pytest.mark.parametrize("variant", FORMATS)
def test_168_layout_structure_matrix(truth: Truth, variant: str):
    filename = f"{truth.number} {truth.zh_title} {truth.en_title}.docx"
    parsed = parse_docx(_save(make_document(truth, variant)), filename)
    actual = [
        {"id": section.id, "zh": section.zh_lines, "en": section.en_lines}
        for section in parsed.sections
    ]
    assert actual == expected_sections(truth)
    assert parsed.unresolved_fragments == []
    assert parsed.classified_fragment_count == parsed.total_fragment_count
    assert parsed.requires_confirmation is False
    assert parsed.confidence >= 0.90
    candidate_kinds = {candidate.layout_kind for candidate in parsed.candidate_layouts}
    if variant == "line_interleaved":
        assert "line_interleaved" in candidate_kinds
    if variant == "table_columns":
        assert "table_columns" in candidate_kinds


def test_reference_manifest_has_twelve_structure_only_records():
    path = Path(__file__).parent / "fixtures" / "docx_lyrics" / "reference_manifest.json"
    records = json.loads(path.read_text(encoding="utf-8"))
    assert len(records) == 12
    assert all(record["content_policy"] == "structure-only" for record in records)
    assert all(record["sources"] for record in records)


def test_numbering_inherited_from_word_list_style_is_visible():
    document = Document()
    document.add_paragraph("第一行", style="List Number")
    document.add_paragraph("第二行", style="List Number")
    visible = [fragment.text for fragment in extract_fragments(document) if fragment.text]
    assert visible == ["1. 第一行", "2. 第二行"]


def test_image_only_input_is_never_automatically_approved():
    document = Document()
    document.add_paragraph("999 圖片詩歌 Image Hymn")
    # A drawing marker without visible text stands in for an unsupported scan.
    drawing = OxmlElement("w:drawing")
    document.add_paragraph()._p.append(drawing)
    parsed = parse_docx(_save(document), "999 圖片詩歌 Image Hymn.docx")
    assert parsed.requires_confirmation is True
    assert any("图片" in item.reason for item in parsed.unresolved_fragments)
