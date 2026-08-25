# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Leo Song
"""Layout-aware, deterministic extraction and parsing for bilingual hymn DOCX files."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
import re
from pathlib import Path
from typing import Iterable, Literal

from docx.document import Document as DocumentType
from docx.oxml.ns import qn

from app.models import (
    DocxLyricsLayoutKind,
    DocxLyricsParseCandidate,
    DocxLyricsSection,
    DocxLyricsUnresolvedFragment,
)
from app.services.chinese_service import contains_chinese


_NUMBERED_TITLE_RE = re.compile(r"^\s*[#№]?\s*(\d+)\s*[.:：\-]?\s+(.+?)\s*$")
_CHORUS_RE = re.compile(
    r"^\s*([\[(（])?\s*(?:副歌|重唱|chorus|cho\.?|refrain|ref\.?)"
    r"\s*([\])）])?\s*(?:(?:[:：])\s*(.*?))?\s*$",
    re.IGNORECASE,
)
_VERSE_RE = re.compile(
    r"^\s*[\[(（]?\s*(?:(?:verse|stanza|v|第)\s*)?"
    r"([0-9]+|[①-⑳])\s*(?:節)?\s*[\])）]?\s*[.:：、\-]?\s*(.*?)\s*$",
    re.IGNORECASE,
)
_AMEN_ONLY_RE = re.compile(r"^\s*(?:amen|a-men|阿[们門])\s*[.!!。！,，;；]*\s*$", re.IGNORECASE)
_AMEN_SUFFIX_RE = re.compile(
    r"^(.*?)\s+((?:amen|a-men)\s*[.!。！,，;；]*)\s*$",
    re.IGNORECASE,
)
_AMEN_SUFFIX_ZH_RE = re.compile(r"^(.*?)(阿[们門]\s*[.!。！,，;；]*)\s*$")
_METADATA_RE = re.compile(
    r"^\s*(?:"
    r"附加.*(?:經文|经文)|"
    r"(?:lyrics?|song\s*text|words?|music|author|composer|translator|tune|meter|"
    r"scripture|source)\s*(?::|：|by\b)|"
    r"(?:copyright|ccli)\b|"
    r"(?:歌詞|歌词|作詞|作词|作曲|譯詞|译词|譯|译|經文|经文|資料|资料)\s*[:：]|"
    r"教會聖詩(?:\s*#\d+)?\s*$|"
    r"hymn(?:'s|s)?\s+for\b"
    r")",
    re.IGNORECASE,
)
_URL_RE = re.compile(r"^(?:https?://|www\.)", re.IGNORECASE)


@dataclass
class SourceFragment:
    id: str
    text: str
    location: str
    source_kind: Literal["paragraph", "table_cell", "textbox", "image"]
    order: int
    table_id: int | None = None
    row: int | None = None
    col: int | None = None
    is_blank: bool = False


@dataclass
class TitleInfo:
    song_number: str
    title_zh: str
    title_en: str
    title_ids: set[str]
    zh_order: int | None
    en_order: int | None


@dataclass
class LanguagePart:
    lines: list[str] = field(default_factory=list)
    fragment_ids: list[str] = field(default_factory=list)
    amen: str | None = None
    amen_fragment_ids: list[str] = field(default_factory=list)


@dataclass
class ParsedLanguage:
    verses: dict[int, LanguagePart] = field(default_factory=dict)
    chorus: LanguagePart | None = None
    explicit_choruses: list[LanguagePart] = field(default_factory=list)
    marker_ids: set[str] = field(default_factory=set)
    chorus_marker_ids: set[str] = field(default_factory=set)
    repeated_chorus: bool = False
    chorus_conflict: bool = False


@dataclass
class CandidateBuild:
    layout_kind: DocxLyricsLayoutKind
    sections: list[DocxLyricsSection]
    marker_ids: set[str]
    grouping_evidence: bool
    repeated_chorus: bool
    missing_declared_chorus: bool
    conflicting_chorus: bool
    fragments: list[SourceFragment]
    title: TitleInfo
    reasons: list[str]


@dataclass
class ParserResult:
    song_number: str
    title_zh: str
    title_en: str
    candidate_layouts: list[DocxLyricsParseCandidate]
    selected: DocxLyricsParseCandidate
    requires_confirmation: bool


class NumberingResolver:
    """Resolve common Word decimal auto-numbering into visible verse markers."""

    def __init__(self, document: DocumentType):
        self.levels: dict[tuple[str, int], tuple[str, int, str]] = {}
        self.counters: dict[tuple[str, int], int] = {}
        self.style_numbering: dict[str, tuple[str, int]] = {}
        try:
            root = document.part.numbering_part.element
        except Exception:
            return
        abstracts: dict[str, object] = {}
        for abstract in root.findall(qn("w:abstractNum")):
            abstract_id = abstract.get(qn("w:abstractNumId"))
            if abstract_id is not None:
                abstracts[abstract_id] = abstract
        for num in root.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            abstract_ref = num.find(qn("w:abstractNumId"))
            if num_id is None or abstract_ref is None:
                continue
            abstract = abstracts.get(abstract_ref.get(qn("w:val"), ""))
            if abstract is None:
                continue
            for level in abstract.findall(qn("w:lvl")):
                ilvl = int(level.get(qn("w:ilvl"), "0"))
                fmt = level.find(qn("w:numFmt"))
                start = level.find(qn("w:start"))
                text = level.find(qn("w:lvlText"))
                self.levels[(num_id, ilvl)] = (
                    fmt.get(qn("w:val"), "decimal") if fmt is not None else "decimal",
                    int(start.get(qn("w:val"), "1")) if start is not None else 1,
                    text.get(qn("w:val"), "%1.") if text is not None else "%1.",
                )
        for style in document.styles.element.findall(qn("w:style")):
            style_id = style.get(qn("w:styleId"))
            num_pr = style.find("./w:pPr/w:numPr", style.nsmap)
            if not style_id or num_pr is None:
                continue
            num_id_el = num_pr.find(qn("w:numId"))
            ilvl_el = num_pr.find(qn("w:ilvl"))
            if num_id_el is not None:
                self.style_numbering[style_id] = (
                    num_id_el.get(qn("w:val"), ""),
                    int(ilvl_el.get(qn("w:val"), "0")) if ilvl_el is not None else 0,
                )

    def prefix(self, paragraph) -> str:
        num_pr = paragraph.find("./w:pPr/w:numPr", paragraph.nsmap)
        if num_pr is not None:
            num_id_el = num_pr.find(qn("w:numId"))
            ilvl_el = num_pr.find(qn("w:ilvl"))
            if num_id_el is None:
                return ""
            num_id = num_id_el.get(qn("w:val"), "")
            ilvl = int(ilvl_el.get(qn("w:val"), "0")) if ilvl_el is not None else 0
        else:
            style = paragraph.find("./w:pPr/w:pStyle", paragraph.nsmap)
            if style is None:
                return ""
            resolved = self.style_numbering.get(style.get(qn("w:val"), ""))
            if resolved is None:
                return ""
            num_id, ilvl = resolved
        fmt, start, template = self.levels.get((num_id, ilvl), ("decimal", 1, "%1."))
        if fmt not in {"decimal", "decimalZero"}:
            return ""
        key = (num_id, ilvl)
        value = self.counters.get(key, start - 1) + 1
        self.counters[key] = value
        for counter_key in list(self.counters):
            if counter_key[0] == num_id and counter_key[1] > ilvl:
                del self.counters[counter_key]
        return template.replace(f"%{ilvl + 1}", str(value)).strip()


def _visible_text(element, *, skip_textboxes: bool = True) -> str:
    pieces: list[str] = []

    def walk(node) -> None:
        if node.tag in {qn("w:del"), qn("w:moveFrom")}:
            return
        if skip_textboxes and node.tag == qn("w:txbxContent"):
            return
        if node.tag == qn("w:t"):
            pieces.append(node.text or "")
            return
        if node.tag == qn("w:tab"):
            pieces.append("\t")
            return
        if node.tag in {qn("w:br"), qn("w:cr")}:
            # A manual line break creates a source line; a page/column break is
            # only a layout boundary and must not terminate the active stanza.
            if node.get(qn("w:type"), "textWrapping") not in {"page", "column"}:
                pieces.append("\n")
            return
        for child in node:
            walk(child)

    walk(element)
    return "".join(pieces)


def _language_guess(text: str) -> Literal["zh", "en", "unknown"]:
    has_zh = contains_chinese(text)
    has_en = bool(re.search(r"[A-Za-z]", text))
    if has_zh and not has_en:
        return "zh"
    if has_en and not has_zh:
        return "en"
    return "unknown"


def extract_fragments(document: DocumentType) -> list[SourceFragment]:
    """Extract visible body content in reading order, including tables/textboxes."""
    fragments: list[SourceFragment] = []
    resolver = NumberingResolver(document)
    sequence = 0
    paragraph_index = 0
    table_index = 0

    def add_text(
        text: str,
        location: str,
        source_kind: Literal["paragraph", "table_cell", "textbox", "image"],
        *,
        table_id: int | None = None,
        row: int | None = None,
        col: int | None = None,
        prefix: str = "",
    ) -> None:
        nonlocal sequence
        raw_lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if not raw_lines:
            raw_lines = [""]
        for line_index, line in enumerate(raw_lines):
            value = line.strip()
            if line_index == 0 and prefix and value:
                value = f"{prefix} {value}".strip()
            fragments.append(
                SourceFragment(
                    id=f"fragment-{sequence + 1}",
                    text=value,
                    location=f"{location}:line:{line_index + 1}",
                    source_kind=source_kind,
                    order=sequence,
                    table_id=table_id,
                    row=row,
                    col=col,
                    is_blank=not value,
                )
            )
            sequence += 1

    def add_paragraph(element, location: str, kind: Literal["paragraph", "table_cell", "textbox"], **kwargs) -> None:
        prefix = resolver.prefix(element) if kind == "paragraph" else ""
        direct = _visible_text(element, skip_textboxes=True)
        has_image = any(
            child.tag in {qn("w:drawing"), qn("w:pict")}
            or child.tag.rsplit("}", 1)[-1] == "imagedata"
            for child in element.iter()
        )
        if direct.strip() or not has_image:
            add_text(direct, location, kind, prefix=prefix, **kwargs)
        if has_image and not direct.strip() and not element.xpath(".//w:txbxContent//w:t"):
            add_text("[图片内容]", location, "image", **kwargs)
        for box_index, text_box in enumerate(element.xpath(".//w:txbxContent")):
            box_paragraphs = list(text_box.iter(qn("w:p")))
            for inner_index, inner_paragraph in enumerate(box_paragraphs):
                add_text(
                    _visible_text(inner_paragraph, skip_textboxes=False),
                    f"{location}:textbox:{box_index + 1}:p:{inner_index + 1}",
                    "textbox",
                    **kwargs,
                )

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_index += 1
            add_paragraph(child, f"body:p:{paragraph_index}", "paragraph")
        elif child.tag == qn("w:tbl"):
            table_index += 1
            for row_index, row in enumerate(child.findall(qn("w:tr"))):
                for col_index, cell in enumerate(row.findall(qn("w:tc"))):
                    paragraphs = cell.xpath("./w:p | ./w:tbl//w:p")
                    if not paragraphs:
                        add_text(
                            "",
                            f"table:{table_index}:row:{row_index + 1}:col:{col_index + 1}",
                            "table_cell",
                            table_id=table_index,
                            row=row_index,
                            col=col_index,
                        )
                    for inner_index, paragraph in enumerate(paragraphs):
                        add_paragraph(
                            paragraph,
                            f"table:{table_index}:row:{row_index + 1}:col:{col_index + 1}:p:{inner_index + 1}",
                            "table_cell",
                            table_id=table_index,
                            row=row_index,
                            col=col_index,
                        )
    return fragments


def _filename_metadata(filename: str) -> tuple[str, str, str]:
    stem = re.sub(r"_v2$", "", Path(filename).stem, flags=re.IGNORECASE).strip()
    match = _NUMBERED_TITLE_RE.match(stem)
    if not match:
        return "", "", ""
    number, remainder = match.groups()
    first_latin = re.search(r"[A-Za-z]", remainder)
    if not first_latin:
        return number, remainder.strip(), ""
    return (
        number,
        remainder[: first_latin.start()].strip(),
        remainder[first_latin.start() :].strip(),
    )


def _normalize(text: str) -> str:
    return re.sub(r"[\W_]+", "", text, flags=re.UNICODE).casefold()


def find_titles(fragments: list[SourceFragment], filename: str) -> TitleInfo:
    number, title_zh, title_en = _filename_metadata(filename)
    title_ids: set[str] = set()
    zh_order: int | None = None
    en_order: int | None = None

    for fragment in fragments:
        if not fragment.text:
            continue
        match = _NUMBERED_TITLE_RE.match(fragment.text)
        if not match:
            continue
        candidate_number, candidate = match.groups()
        if contains_chinese(candidate):
            number = candidate_number
            first_latin = re.search(r"[A-Za-z]", candidate)
            if first_latin:
                title_zh = candidate[: first_latin.start()].strip() or title_zh
                title_en = candidate[first_latin.start() :].strip() or title_en
                en_order = fragment.order
            else:
                title_zh = candidate.strip() or title_zh
            zh_order = fragment.order
            title_ids.add(fragment.id)
            break

    if zh_order is None and title_zh:
        wanted_zh = _normalize(title_zh)
        first_marker_order = next(
            (
                fragment.order
                for fragment in fragments
                if fragment.text and _marker(fragment.text)[0]
            ),
            None,
        )
        for fragment in fragments:
            if first_marker_order is not None and fragment.order >= first_marker_order:
                break
            if fragment.text and _normalize(fragment.text) == wanted_zh:
                zh_order = fragment.order
                title_ids.add(fragment.id)
                break

    search_start = (zh_order + 1) if zh_order is not None else 0
    wanted_en = _normalize(title_en)
    for fragment in fragments:
        if fragment.order < search_start or not fragment.text:
            continue
        match = _NUMBERED_TITLE_RE.match(fragment.text)
        if match:
            candidate_number, candidate = match.groups()
            if not contains_chinese(candidate) and (not number or candidate_number == number):
                number = number or candidate_number
                title_en = candidate.strip()
                en_order = fragment.order
                title_ids.add(fragment.id)
                break
        if wanted_en and _normalize(fragment.text) == wanted_en:
            title_en = fragment.text.strip()
            en_order = fragment.order
            title_ids.add(fragment.id)
            break

    return TitleInfo(number, title_zh, title_en, title_ids, zh_order, en_order)


def _marker(text: str) -> tuple[str | None, int | None, str, bool]:
    chorus = _CHORUS_RE.match(text)
    if chorus:
        remainder = (chorus.group(3) or "").strip()
        is_reference = not remainder and bool(chorus.group(1) and chorus.group(2))
        return "chorus", None, remainder, is_reference
    verse = _VERSE_RE.match(text)
    if verse:
        raw_number, remainder = verse.groups()
        if raw_number.isdigit():
            number = int(raw_number)
        else:
            number = ord(raw_number) - ord("①") + 1
        # A bare lyric line beginning with a number must have a delimiter or label.
        if remainder and not re.match(r"^\s*(?:verse|stanza|v|第)", text, re.IGNORECASE):
            prefix = text[: text.find(raw_number) + len(raw_number)]
            after = text[len(prefix) :]
            if not re.match(r"^\s*[\])）.:：、\-]", after):
                return None, None, text, False
        return "verse", number, remainder.strip(), False
    return None, None, text, False


def _groups(fragments: Iterable[SourceFragment]) -> list[list[SourceFragment]]:
    groups: list[list[SourceFragment]] = []
    current: list[SourceFragment] = []
    for fragment in fragments:
        if fragment.is_blank:
            if current:
                groups.append(current)
                current = []
        else:
            current.append(fragment)
    if current:
        groups.append(current)
    return groups


def _split_amen(
    lines: list[str], ids: list[str]
) -> tuple[list[str], list[str], str | None, list[str]]:
    if not lines:
        return lines, ids, None, []
    last = lines[-1]
    if _AMEN_ONLY_RE.match(last):
        return lines[:-1], ids[:-1], last, ids[-1:]
    suffix = _AMEN_SUFFIX_RE.match(last)
    if not suffix:
        suffix = _AMEN_SUFFIX_ZH_RE.match(last)
    if suffix and suffix.group(1).strip():
        return [*lines[:-1], suffix.group(1).strip()], ids, suffix.group(2), ids[-1:]
    return lines, ids, None, []


def _common_suffix(parts: list[LanguagePart]) -> tuple[list[str], int]:
    if len(parts) < 2:
        return [], 0
    core_parts: list[list[str]] = []
    normalized_parts: list[list[str]] = []
    for part in parts:
        core, _, _, _ = _split_amen(part.lines, part.fragment_ids)
        core_parts.append(core)
        normalized_parts.append([_normalize(line) for line in core])
    # A repeated refrain must leave actual verse text in every stanza. Treating
    # a set of identical short verses as an all-chorus song is unsafe.
    max_length = min(max(0, len(lines) - 2) for lines in normalized_parts)
    length = 0
    for offset in range(1, max_length + 1):
        values = {lines[-offset] for lines in normalized_parts}
        if len(values) != 1 or not next(iter(values)):
            break
        length = offset
    if length == 0:
        return [], 0
    suffix = core_parts[0][-length:]
    if length < 2 and sum(len(line) for line in suffix) < 20:
        return [], 0
    return suffix, length


def _finish_explicit_choruses(parsed: ParsedLanguage) -> None:
    occurrences: list[LanguagePart] = []
    ending_only: list[LanguagePart] = []
    for part in parsed.explicit_choruses:
        part.lines, part.fragment_ids, detected_amen, amen_ids = _split_amen(
            part.lines, part.fragment_ids
        )
        part.amen = part.amen or detected_amen
        part.amen_fragment_ids = [*part.amen_fragment_ids, *amen_ids]
        if part.lines:
            occurrences.append(part)
        elif part.amen:
            ending_only.append(part)
    if not occurrences:
        return

    canonical = occurrences[0]
    canonical_signature = tuple(_normalize(line) for line in canonical.lines)
    for occurrence in occurrences[1:]:
        signature = tuple(_normalize(line) for line in occurrence.lines)
        if signature != canonical_signature:
            parsed.chorus_conflict = True
            continue
        canonical.fragment_ids.extend(occurrence.fragment_ids)
        canonical.amen_fragment_ids.extend(occurrence.amen_fragment_ids)
        if occurrence.amen:
            canonical.amen = occurrence.amen
        parsed.repeated_chorus = True
    for ending in ending_only:
        canonical.amen_fragment_ids.extend(ending.amen_fragment_ids)
        canonical.amen = ending.amen
    parsed.chorus = canonical


def _extract_repeated_chorus(parsed: ParsedLanguage) -> None:
    if parsed.chorus or len(parsed.verses) < 2:
        return
    ordered = [parsed.verses[number] for number in sorted(parsed.verses)]
    suffix, length = _common_suffix(ordered)
    if not suffix:
        return
    chorus_ids: list[str] = []
    amen_fragment_ids: list[str] = []
    amen: str | None = None
    for part in ordered:
        core, core_ids, detected_amen, amen_ids = _split_amen(part.lines, part.fragment_ids)
        if detected_amen:
            amen = detected_amen
            amen_fragment_ids.extend(amen_ids)
        chorus_ids.extend(core_ids[-length:])
        part.lines = core[:-length]
        part.fragment_ids = core_ids[:-length]
    parsed.chorus = LanguagePart(
        suffix,
        chorus_ids,
        amen,
        amen_fragment_ids,
    )
    parsed.repeated_chorus = True


def _matches_language(fragment: SourceFragment, language: str) -> bool:
    guess = _language_guess(fragment.text)
    return guess == language or guess == "unknown"


def parse_single_language(fragments: list[SourceFragment], language: str) -> ParsedLanguage:
    fragments = [
        fragment
        for fragment in fragments
        if fragment.is_blank or not _METADATA_RE.match(fragment.text)
    ]
    parsed = ParsedLanguage()
    explicit = any(_marker(fragment.text)[0] for fragment in fragments if fragment.text)
    if explicit:
        current_kind = "verse"
        current_number = 1
        current_chorus: LanguagePart | None = None
        after_chorus_blank = False
        for fragment in fragments:
            if fragment.is_blank:
                if current_kind == "chorus" and current_chorus and current_chorus.lines:
                    after_chorus_blank = True
                continue
            kind, number, remainder, is_reference = _marker(fragment.text)
            if kind:
                parsed.marker_ids.add(fragment.id)
                if kind == "chorus":
                    parsed.chorus_marker_ids.add(fragment.id)
                if is_reference:
                    continue
                if kind == "verse":
                    current_kind = "verse"
                    current_number = number or current_number
                    current_chorus = None
                    after_chorus_blank = False
                else:
                    current_kind = "chorus"
                    current_chorus = LanguagePart()
                    parsed.explicit_choruses.append(current_chorus)
                    after_chorus_blank = False
                if not remainder:
                    continue
                line = remainder
            else:
                if current_kind == "chorus" and after_chorus_blank:
                    current_kind = "verse"
                    current_number = max(parsed.verses, default=0) + 1
                    after_chorus_blank = False
                line = fragment.text
            if not _matches_language(fragment, language):
                continue
            if current_kind == "chorus":
                if current_chorus is None:
                    current_chorus = LanguagePart()
                    parsed.explicit_choruses.append(current_chorus)
                current_chorus.lines.append(line)
                current_chorus.fragment_ids.append(fragment.id)
            else:
                part = parsed.verses.setdefault(current_number, LanguagePart())
                part.lines.append(line)
                part.fragment_ids.append(fragment.id)
    else:
        next_is_chorus = False
        for group in _groups(fragments):
            matching = [item for item in group if _matches_language(item, language)]
            if not matching:
                continue
            marker_index = next(
                (index for index, item in enumerate(matching) if _marker(item.text)[0] == "chorus"),
                None,
            )
            if marker_index is not None:
                before = matching[:marker_index]
                marker_fragment = matching[marker_index]
                parsed.marker_ids.add(marker_fragment.id)
                if before:
                    number = len(parsed.verses) + 1
                    parsed.verses[number] = LanguagePart(
                        [item.text for item in before], [item.id for item in before]
                    )
                remainder = _marker(marker_fragment.text)[2]
                after = matching[marker_index + 1 :]
                chorus_lines = ([remainder] if remainder else []) + [item.text for item in after]
                chorus_ids = ([marker_fragment.id] if remainder else []) + [item.id for item in after]
                if chorus_lines:
                    parsed.chorus = LanguagePart(chorus_lines, chorus_ids)
                    next_is_chorus = False
                else:
                    next_is_chorus = True
                continue
            if next_is_chorus:
                parsed.chorus = LanguagePart(
                    [item.text for item in matching], [item.id for item in matching]
                )
                next_is_chorus = False
                continue
            number = len(parsed.verses) + 1
            parsed.verses[number] = LanguagePart(
                [item.text for item in matching], [item.id for item in matching]
            )
    _finish_explicit_choruses(parsed)
    _extract_repeated_chorus(parsed)
    for part in [*parsed.verses.values(), *([parsed.chorus] if parsed.chorus else [])]:
        if part is None:
            continue
        part.lines, part.fragment_ids, detected_amen, amen_ids = _split_amen(
            part.lines, part.fragment_ids
        )
        part.amen = part.amen or detected_amen
        part.amen_fragment_ids = part.amen_fragment_ids or amen_ids
    return parsed


def _combine_languages(zh: ParsedLanguage, en: ParsedLanguage) -> list[DocxLyricsSection]:
    sections: list[DocxLyricsSection] = []
    for number in sorted(set(zh.verses) | set(en.verses)):
        zh_part = zh.verses.get(number, LanguagePart())
        en_part = en.verses.get(number, LanguagePart())
        sections.append(
            DocxLyricsSection(
                id=f"verse-{number}",
                kind="verse",
                number=number,
                zh_lines=zh_part.lines,
                en_lines=en_part.lines,
                source_fragment_ids=[
                    *zh_part.fragment_ids,
                    *zh_part.amen_fragment_ids,
                    *en_part.fragment_ids,
                    *en_part.amen_fragment_ids,
                ],
                amen_zh=zh_part.amen,
                amen_en=en_part.amen,
            )
        )
    if zh.chorus or en.chorus:
        zh_part = zh.chorus or LanguagePart()
        en_part = en.chorus or LanguagePart()
        sections.append(
            DocxLyricsSection(
                id="chorus",
                kind="chorus",
                zh_lines=zh_part.lines,
                en_lines=en_part.lines,
                source_fragment_ids=[
                    *zh_part.fragment_ids,
                    *zh_part.amen_fragment_ids,
                    *en_part.fragment_ids,
                    *en_part.amen_fragment_ids,
                ],
                amen_zh=zh_part.amen,
                amen_en=en_part.amen,
            )
        )
    return sections


def _content_after_titles(fragments: list[SourceFragment], title: TitleInfo) -> list[SourceFragment]:
    orders = [order for order in (title.zh_order, title.en_order) if order is not None]
    start = min(orders) + 1 if orders else 0
    return [fragment for fragment in fragments if fragment.order >= start and fragment.id not in title.title_ids]


def build_language_blocks(fragments: list[SourceFragment], title: TitleInfo) -> CandidateBuild | None:
    if title.zh_order is None or title.en_order is None or title.en_order <= title.zh_order:
        return None
    zh_fragments = [
        fragment
        for fragment in fragments
        if title.zh_order < fragment.order < title.en_order
        and fragment.id not in title.title_ids
    ]
    en_fragments = [
        fragment
        for fragment in fragments
        if fragment.order > title.en_order and fragment.id not in title.title_ids
    ]
    if not any(item.text for item in zh_fragments) or not any(item.text for item in en_fragments):
        return None
    zh = parse_single_language(zh_fragments, "zh")
    en = parse_single_language(en_fragments, "en")
    sections = _combine_languages(zh, en)
    if not sections:
        return None
    return CandidateBuild(
        "language_blocks",
        sections,
        zh.marker_ids | en.marker_ids,
        True,
        zh.repeated_chorus or en.repeated_chorus,
        bool(
            (zh.chorus_marker_ids and not zh.chorus)
            or (en.chorus_marker_ids and not en.chorus)
        ),
        zh.chorus_conflict or en.chorus_conflict,
        fragments,
        title,
        ["识别为中文歌词块后接英文歌词块"],
    )


def _mixed_groups_candidate(
    fragments: list[SourceFragment], title: TitleInfo, layout: DocxLyricsLayoutKind
) -> CandidateBuild | None:
    content = _content_after_titles(fragments, title)
    content = [
        fragment
        for fragment in content
        if fragment.is_blank or not _METADATA_RE.match(fragment.text)
    ]
    parsed_groups: list[tuple[LanguagePart, LanguagePart]] = []
    marker_ids: set[str] = set()
    chorus_marker_ids: set[str] = set()
    current_zh = ParsedLanguage()
    current_en = ParsedLanguage()
    active_zh_chorus: LanguagePart | None = None
    active_en_chorus: LanguagePart | None = None
    has_explicit = any(_marker(item.text)[0] for item in content if item.text)

    if has_explicit:
        current_kind = "verse"
        current_number = 1
        for fragment in content:
            if fragment.is_blank:
                continue
            kind, number, remainder, is_reference = _marker(fragment.text)
            if kind:
                marker_ids.add(fragment.id)
                if kind == "chorus":
                    chorus_marker_ids.add(fragment.id)
                if is_reference:
                    continue
                current_kind = kind
                if kind == "verse":
                    current_number = number or current_number
                    active_zh_chorus = None
                    active_en_chorus = None
                else:
                    active_zh_chorus = None
                    active_en_chorus = None
                    marker_language = _language_guess(fragment.text)
                    if marker_language != "en":
                        active_zh_chorus = LanguagePart()
                        current_zh.explicit_choruses.append(active_zh_chorus)
                    if marker_language != "zh":
                        active_en_chorus = LanguagePart()
                        current_en.explicit_choruses.append(active_en_chorus)
                if not remainder:
                    continue
                line = remainder
            else:
                line = fragment.text
            language = _language_guess(line)
            if language not in {"zh", "en"}:
                continue
            parsed = current_zh if language == "zh" else current_en
            if current_kind == "chorus":
                active = active_zh_chorus if language == "zh" else active_en_chorus
                if active is None:
                    active = LanguagePart()
                    parsed.explicit_choruses.append(active)
                    if language == "zh":
                        active_zh_chorus = active
                    else:
                        active_en_chorus = active
                active.lines.append(line)
                active.fragment_ids.append(fragment.id)
            else:
                part = parsed.verses.setdefault(current_number, LanguagePart())
                part.lines.append(line)
                part.fragment_ids.append(fragment.id)
    else:
        for group in _groups(content):
            zh_items = [item for item in group if _language_guess(item.text) == "zh"]
            en_items = [item for item in group if _language_guess(item.text) == "en"]
            if zh_items and en_items:
                parsed_groups.append(
                    (
                        LanguagePart([item.text for item in zh_items], [item.id for item in zh_items]),
                        LanguagePart([item.text for item in en_items], [item.id for item in en_items]),
                    )
                )
        if not parsed_groups:
            return None
        for index, (zh_part, en_part) in enumerate(parsed_groups, start=1):
            current_zh.verses[index] = zh_part
            current_en.verses[index] = en_part
        _extract_repeated_chorus(current_zh)
        _extract_repeated_chorus(current_en)

    _finish_explicit_choruses(current_zh)
    _finish_explicit_choruses(current_en)

    # Explicit verse labels may coexist with an unlabelled refrain repeated at
    # the end of every stanza. The repeated-suffix pass is therefore required
    # for both explicit-marker and blank-group layouts.
    _extract_repeated_chorus(current_zh)
    _extract_repeated_chorus(current_en)

    for parsed in (current_zh, current_en):
        for part in [*parsed.verses.values(), *([parsed.chorus] if parsed.chorus else [])]:
            if part is None:
                continue
            part.lines, part.fragment_ids, detected_amen, amen_ids = _split_amen(
                part.lines, part.fragment_ids
            )
            part.amen = part.amen or detected_amen
            part.amen_fragment_ids = part.amen_fragment_ids or amen_ids
    sections = _combine_languages(current_zh, current_en)
    if not sections:
        return None
    return CandidateBuild(
        layout,
        sections,
        marker_ids,
        bool(parsed_groups) or has_explicit,
        current_zh.repeated_chorus or current_en.repeated_chorus,
        bool(chorus_marker_ids) and not (current_zh.chorus and current_en.chorus),
        current_zh.chorus_conflict or current_en.chorus_conflict,
        fragments,
        title,
        [
            "识别为逐行中英配对"
            if layout == "line_interleaved"
            else "识别为逐节中英配对"
        ],
    )


def _alternation_ratio(fragments: list[SourceFragment], title: TitleInfo) -> float:
    languages = [
        _language_guess(item.text)
        for item in _content_after_titles(fragments, title)
        if item.text and _language_guess(item.text) in {"zh", "en"}
    ]
    if len(languages) < 2:
        return 0.0
    changes = sum(left != right for left, right in zip(languages, languages[1:]))
    return changes / (len(languages) - 1)


def build_table_columns(fragments: list[SourceFragment], title: TitleInfo) -> CandidateBuild | None:
    table_fragments = [item for item in fragments if item.table_id is not None]
    if not table_fragments:
        return None
    by_table: dict[int, list[SourceFragment]] = defaultdict(list)
    for fragment in table_fragments:
        by_table[fragment.table_id or 0].append(fragment)
    best: tuple[list[SourceFragment], list[SourceFragment]] | None = None
    for items in by_table.values():
        columns: dict[int, list[SourceFragment]] = defaultdict(list)
        for item in items:
            columns[item.col or 0].append(item)
        if len(columns) < 2:
            continue
        ranked = sorted(columns)
        left, right = columns[ranked[0]], columns[ranked[1]]
        left_zh = sum(_language_guess(item.text) == "zh" for item in left)
        right_en = sum(_language_guess(item.text) == "en" for item in right)
        reverse = (
            sum(_language_guess(item.text) == "en" for item in left)
            + sum(_language_guess(item.text) == "zh" for item in right)
        )
        if reverse > left_zh + right_en:
            left, right = right, left
            left_zh = sum(_language_guess(item.text) == "zh" for item in left)
            right_en = sum(_language_guess(item.text) == "en" for item in right)
        if left_zh and right_en:
            best = (left, right)
            break
    if best is None:
        return None
    zh_fragments, en_fragments = best

    def with_row_breaks(items: list[SourceFragment]) -> list[SourceFragment]:
        result: list[SourceFragment] = []
        previous_row: int | None = None
        for item in sorted(items, key=lambda value: (value.row or 0, value.order)):
            if previous_row is not None and item.row != previous_row:
                result.append(
                    SourceFragment(
                        id=f"synthetic-table-break-{item.table_id}-{item.row}-{item.col}",
                        text="",
                        location=item.location,
                        source_kind="table_cell",
                        order=item.order,
                        table_id=item.table_id,
                        row=item.row,
                        col=item.col,
                        is_blank=True,
                    )
                )
            result.append(item)
            previous_row = item.row
        return result

    zh = parse_single_language(
        [item for item in with_row_breaks(zh_fragments) if item.id not in title.title_ids], "zh"
    )
    en = parse_single_language(
        [item for item in with_row_breaks(en_fragments) if item.id not in title.title_ids], "en"
    )
    sections = _combine_languages(zh, en)
    if not sections:
        return None
    return CandidateBuild(
        "table_columns",
        sections,
        zh.marker_ids | en.marker_ids,
        True,
        zh.repeated_chorus or en.repeated_chorus,
        bool(
            (zh.chorus_marker_ids and not zh.chorus)
            or (en.chorus_marker_ids and not en.chorus)
        ),
        zh.chorus_conflict or en.chorus_conflict,
        fragments,
        title,
        ["识别为 Word 表格中的中英双栏"],
    )


def _candidate_signature(sections: list[DocxLyricsSection]) -> tuple:
    return tuple(
        (
            section.kind,
            section.number,
            tuple(_normalize(line) for line in section.zh_lines),
            tuple(_normalize(line) for line in section.en_lines),
            _normalize(section.amen_zh or ""),
            _normalize(section.amen_en or ""),
        )
        for section in sections
    )


def _metadata_ids(fragments: list[SourceFragment], title: TitleInfo) -> set[str]:
    ids = set(title.title_ids)
    normalized_titles = {_normalize(title.title_zh), _normalize(title.title_en)} - {""}
    for fragment in fragments:
        text = fragment.text.strip()
        if not text:
            continue
        if _METADATA_RE.match(text) or _URL_RE.match(text) or _normalize(text) in normalized_titles:
            ids.add(fragment.id)
    return ids


def _finalize_candidate(build: CandidateBuild, index: int) -> DocxLyricsParseCandidate:
    lyric_ids = {
        fragment_id
        for section in build.sections
        for fragment_id in section.source_fragment_ids
        if not fragment_id.startswith("synthetic-")
    }
    classified = lyric_ids | build.marker_ids | _metadata_ids(build.fragments, build.title)
    nonblank = [fragment for fragment in build.fragments if fragment.text]
    unresolved: list[DocxLyricsUnresolvedFragment] = []
    for fragment in nonblank:
        if fragment.id in classified:
            continue
        reason = (
            "图片内文字无法在本功能中识别"
            if fragment.source_kind == "image"
            else "该文字无法安全归入标题、歌词或备注"
        )
        unresolved.append(
            DocxLyricsUnresolvedFragment(
                id=fragment.id,
                text=fragment.text,
                location=fragment.location,
                language_guess=_language_guess(fragment.text),
                reason=reason,
            )
        )
    total = len(nonblank)
    classified_count = total - len(unresolved)
    coverage = classified_count / total if total else 0.0
    title_score = 0.15 if all(
        value.strip() for value in (build.title.song_number, build.title.title_zh, build.title.title_en)
    ) else 0.05
    verses = [section for section in build.sections if section.kind == "verse"]
    paired = sum(bool(section.zh_lines) and bool(section.en_lines) for section in verses)
    chorus = next((section for section in build.sections if section.kind == "chorus"), None)
    paired_score = 0.25 * (paired / len(verses)) if verses else 0.0
    marker_score = 0.20 if build.marker_ids else (0.16 if build.grouping_evidence else 0.0)
    if chorus is None:
        chorus_score = 0.0 if build.missing_declared_chorus else 0.15
    elif chorus.zh_lines and chorus.en_lines and (build.marker_ids or build.repeated_chorus):
        chorus_score = 0.15
    else:
        chorus_score = 0.07
    pairing_ratio = paired / len(verses) if verses else 0.0
    confidence = min(
        1.0,
        title_score + paired_score + marker_score + chorus_score + 0.25 * coverage,
    )
    # A candidate with an unmatched language is unsafe even when it happened to
    # account for every source fragment. Penalize that ambiguity explicitly so
    # it cannot crowd a structurally complete candidate at the 0.15 margin gate.
    confidence = max(0.0, confidence - 0.55 * (1.0 - pairing_ratio))
    if unresolved:
        # Unaccounted source text is a hard safety signal. Keep such a
        # candidate below both the automatic-pass threshold and the margin that
        # could crowd a fully classified alternative.
        confidence = min(confidence, 0.84)
    if build.missing_declared_chorus:
        confidence = min(confidence, 0.84)
    if build.conflicting_chorus:
        confidence = min(confidence, 0.84)
    reasons = [*build.reasons]
    reasons.append(f"源文归类 {classified_count}/{total}")
    if build.marker_ids:
        reasons.append("存在明确 Verse/Chorus 标记")
    if build.repeated_chorus:
        reasons.append("识别到多节重复的副歌尾段")
    if build.missing_declared_chorus:
        reasons.append("发现副歌标签或引用，但没有识别到完整双语副歌正文")
    if build.conflicting_chorus:
        reasons.append("发现多个文本不同的副歌，需人工确认")
    sequence: list[str] = []
    for verse in verses:
        sequence.append(verse.id)
        if chorus:
            sequence.append(chorus.id)
    return DocxLyricsParseCandidate(
        id=f"candidate-{index}",
        layout_kind=build.layout_kind,
        confidence=round(confidence, 3),
        reasons=reasons,
        sections=build.sections,
        sequence=sequence,
        unresolved_fragments=unresolved,
        classified_fragment_count=classified_count,
        total_fragment_count=total,
    )


def parse_document(document: DocumentType, filename: str) -> ParserResult:
    fragments = extract_fragments(document)
    title = find_titles(fragments, filename)
    builds: list[CandidateBuild] = []
    language_blocks = build_language_blocks(fragments, title)
    if language_blocks:
        builds.append(language_blocks)
    stanza = _mixed_groups_candidate(fragments, title, "stanza_interleaved")
    if stanza:
        builds.append(stanza)
    if _alternation_ratio(fragments, title) >= 0.55:
        line = _mixed_groups_candidate(fragments, title, "line_interleaved")
        if line:
            line.reasons.append("中英行交替率较高")
            builds.append(line)
    table = build_table_columns(fragments, title)
    if table:
        builds.append(table)
    if not builds:
        empty = DocxLyricsSection(id="verse-1", kind="verse", number=1)
        builds.append(
            CandidateBuild(
                "language_blocks", [empty], set(), False, False, False, False,
                fragments, title,
                ["未找到可靠的双语歌词布局"],
            )
        )

    candidates = [
        _finalize_candidate(build, index + 1)
        for index, build in enumerate(builds)
    ]
    candidates.sort(key=lambda candidate: candidate.confidence, reverse=True)
    selected = candidates[0]
    selected_signature = _candidate_signature(selected.sections)
    structural_alternatives = [
        candidate
        for candidate in candidates[1:]
        if _candidate_signature(candidate.sections) != selected_signature
    ]
    second_score = structural_alternatives[0].confidence if structural_alternatives else 0.0
    margin = selected.confidence - second_score
    sections_complete = bool(selected.sections) and all(
        section.zh_lines and section.en_lines for section in selected.sections
    )
    verse_numbers = [
        section.number for section in selected.sections if section.kind == "verse"
    ]
    consecutive = verse_numbers == list(range(1, len(verse_numbers) + 1))
    requires_confirmation = not (
        selected.confidence >= 0.90
        and margin >= 0.15
        and not selected.unresolved_fragments
        and sections_complete
        and consecutive
        and all((title.song_number, title.title_zh, title.title_en))
    )
    return ParserResult(
        song_number=title.song_number,
        title_zh=title.title_zh,
        title_en=title.title_en,
        candidate_layouts=candidates,
        selected=selected,
        requires_confirmation=requires_confirmation,
    )
