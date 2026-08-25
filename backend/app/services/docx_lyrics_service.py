# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Leo Song
"""Deterministic bilingual hymn DOCX import and copy-friendly export."""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

from app.models import (
    DocxLyricsExportRequest,
    DocxLyricsImportResponse,
    DocxLyricsSection,
    DocxLyricsWarning,
)
from app.services import docx_lyrics_parser


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_COLLECTION_EN = "Hymn’s for God’s People"
_V2_SUFFIX_RE = re.compile(r"_v2$", re.IGNORECASE)


class DocxLyricsError(ValueError):
    """A user-correctable DOCX import/export error."""


def _safe_source_name(filename: str) -> str:
    name = Path(filename or "lyrics.docx").name
    if Path(name).suffix.lower() != ".docx":
        raise DocxLyricsError("仅支持 .docx 文件")
    if _V2_SUFFIX_RE.search(Path(name).stem):
        raise DocxLyricsError("请选择原始 DOCX，不要再次导入已经生成的 _v2 文件")
    return name


def output_filename(filename: str) -> str:
    name = Path(filename or "lyrics.docx").name
    stem = _V2_SUFFIX_RE.sub("", Path(name).stem).rstrip()
    return f"{stem or 'lyrics'}_v2.docx"


def validate_content(
    song_number: str,
    title_zh: str,
    title_en: str,
    sections: list[DocxLyricsSection],
) -> list[DocxLyricsWarning]:
    warnings: list[DocxLyricsWarning] = []

    def error(code: str, message: str, section_id: str | None = None) -> None:
        warnings.append(
            DocxLyricsWarning(
                code=code, message=message, severity="error", section_id=section_id
            )
        )

    if not song_number.strip():
        error("missing_song_number", "未识别到诗歌编号")
    if not title_zh.strip():
        error("missing_title_zh", "未识别到中文歌名")
    if not title_en.strip():
        error("missing_title_en", "未识别到英文歌名")

    verses = [section for section in sections if section.kind == "verse"]
    choruses = [section for section in sections if section.kind == "chorus"]
    if not verses:
        error("missing_verses", "未识别到 Verse 歌词")
    verse_numbers = [section.number for section in verses]
    if verse_numbers != list(range(1, len(verse_numbers) + 1)):
        error("invalid_verse_order", "Verse 编号必须从 1 开始连续排列")
    if len(choruses) > 1:
        error("multiple_choruses", "传统圣诗模式只支持一个副歌")
    for section in sections:
        label = "副歌" if section.kind == "chorus" else f"第 {section.number} 节"
        if not any(line.strip() for line in section.zh_lines):
            error("missing_section_zh", f"{label}缺少中文歌词", section.id)
        if not any(line.strip() for line in section.en_lines):
            error("missing_section_en", f"{label}缺少英文歌词", section.id)
    return warnings


def expanded_sections(sections: list[DocxLyricsSection]) -> list[DocxLyricsSection]:
    verses = sorted(
        (section for section in sections if section.kind == "verse"),
        key=lambda section: section.number or 0,
    )
    chorus = next((section for section in sections if section.kind == "chorus"), None)
    expanded: list[DocxLyricsSection] = []
    for verse in verses:
        expanded.append(verse)
        if chorus:
            expanded.append(chorus)
    return expanded


def _final_endings(
    sections: list[DocxLyricsSection],
) -> tuple[str | None, str | None]:
    ending_zh = next(
        (section.amen_zh for section in reversed(sections) if section.amen_zh),
        None,
    )
    ending_en = next(
        (section.amen_en for section in reversed(sections) if section.amen_en),
        None,
    )
    return ending_zh, ending_en


def format_lyrics(
    sections: list[DocxLyricsSection],
) -> tuple[str, str, str, list[str]]:
    expanded = expanded_sections(sections)
    ending_zh, ending_en = _final_endings(sections)
    primary_blocks: list[str] = []
    secondary_blocks: list[str] = []
    combined_blocks: list[str] = []
    for index, section in enumerate(expanded):
        is_final = index == len(expanded) - 1
        zh_lines = [*section.zh_lines]
        en_lines = [*section.en_lines]
        if is_final and ending_zh:
            zh_lines.append(ending_zh)
        if is_final and ending_en:
            en_lines.append(ending_en)
        primary_blocks.append("\n".join(zh_lines).strip())
        secondary_blocks.append("\n".join(en_lines).strip())
        combined_blocks.append("\n".join([*zh_lines, *en_lines]).strip())
    sequence = [section.id for section in expanded]
    return (
        "\n\n".join(block for block in primary_blocks if block),
        "\n\n".join(block for block in secondary_blocks if block),
        "\n\n".join(block for block in combined_blocks if block),
        sequence,
    )


def parse_docx(data: bytes, filename: str) -> DocxLyricsImportResponse:
    source_filename = _safe_source_name(filename)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocxLyricsError("无法读取这个 DOCX；文件可能已经损坏") from exc

    parser_result = docx_lyrics_parser.parse_document(document, source_filename)
    if parser_result.selected.total_fragment_count == 0:
        raise DocxLyricsError("DOCX 中没有可读取的文字")
    song_number = parser_result.song_number
    title_zh = parser_result.title_zh
    title_en = parser_result.title_en
    selected = parser_result.selected
    sections = selected.sections
    warnings = validate_content(song_number, title_zh, title_en, sections)
    if selected.unresolved_fragments:
        warnings.append(
            DocxLyricsWarning(
                code="unresolved_fragments",
                message=f"有 {len(selected.unresolved_fragments)} 处原文尚未归类",
                severity="error",
            )
        )
    if parser_result.requires_confirmation:
        warnings.append(
            DocxLyricsWarning(
                code="manual_confirmation_required",
                message="分段存在歧义，请人工检查并确认",
                severity="error",
            )
        )
    primary, secondary, combined, sequence = format_lyrics(sections)
    collection_zh = f"教會聖詩 #{song_number}" if song_number else "教會聖詩"
    return DocxLyricsImportResponse(
        source_filename=source_filename,
        output_filename=output_filename(source_filename),
        song_number=song_number,
        title_zh=title_zh,
        title_en=title_en,
        collection_zh=collection_zh,
        collection_en=DEFAULT_COLLECTION_EN,
        sections=sections,
        sequence=sequence,
        primary_lyrics=primary,
        secondary_lyrics=secondary,
        combined_lyrics=combined,
        warnings=warnings,
        has_blocking_errors=any(w.severity == "error" for w in warnings),
        layout_kind=selected.layout_kind,
        confidence=selected.confidence,
        requires_confirmation=parser_result.requires_confirmation,
        review_confirmed=False,
        candidate_layouts=parser_result.candidate_layouts,
        unresolved_fragments=selected.unresolved_fragments,
        classified_fragment_count=selected.classified_fragment_count,
        total_fragment_count=selected.total_fragment_count,
    )


def _style_run(run, language: str) -> None:
    run.font.size = Pt(12)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    if language == "zh":
        # Keep the OOXML declaration faithful to the requested Windows family.
        # The font table below maps it to macOS's installed Songti TC family.
        run.font.name = "宋体"
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "宋体")
    else:
        run.font.name = "Times New Roman"
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Times New Roman")


def _add_line(document: Document, text: str = "", language: str = "en") -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if text:
        run = paragraph.add_run(text)
        _style_run(run, language)


def _add_songti_font_alias(document: Document) -> None:
    """Map the requested 宋体 family to macOS's installed Songti TC."""
    font_part = next(
        (
            part
            for part in document.part.package.parts
            if str(part.partname) == "/word/fontTable.xml"
        ),
        None,
    )
    if font_part is None:
        return
    root = parse_xml(font_part.blob)
    if any(
        child.get(qn("w:name")) == "宋体"
        for child in root.findall(qn("w:font"))
    ):
        return
    font = OxmlElement("w:font")
    font.set(qn("w:name"), "宋体")
    alt_name = OxmlElement("w:altName")
    alt_name.set(qn("w:val"), "Songti TC")
    font.append(alt_name)
    family = OxmlElement("w:family")
    family.set(qn("w:val"), "roman")
    font.append(family)
    root.append(font)
    font_part._blob = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def build_docx(request: DocxLyricsExportRequest) -> tuple[bytes, str]:
    warnings = validate_content(
        request.song_number, request.title_zh, request.title_en, request.sections
    )
    errors = [warning.message for warning in warnings if warning.severity == "error"]
    if errors:
        raise DocxLyricsError("；".join(errors))

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    _add_line(document, request.title_zh.strip(), "zh")
    _add_line(document, request.title_en.strip(), "en")
    _add_line(document)
    _add_line(document, request.collection_zh.strip(), "zh")
    _add_line(document, request.collection_en.strip(), "en")
    _add_line(document)
    _add_line(document)

    expanded = expanded_sections(request.sections)
    ending_zh, ending_en = _final_endings(request.sections)
    for section_index, lyrics_section in enumerate(expanded):
        is_final = section_index == len(expanded) - 1
        zh_lines = [*lyrics_section.zh_lines]
        en_lines = [*lyrics_section.en_lines]
        if is_final and ending_zh:
            zh_lines.append(ending_zh)
        if is_final and ending_en:
            en_lines.append(ending_en)
        for line in zh_lines:
            _add_line(document, line, "zh")
        for line in en_lines:
            _add_line(document, line, "en")
        if section_index < len(expanded) - 1:
            _add_line(document)

    _add_songti_font_alias(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), output_filename(request.source_filename)
